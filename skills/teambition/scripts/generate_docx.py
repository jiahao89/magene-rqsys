"""
Word 周报文档生成器 — 读取 generate.py 产出的 JSON，写入模板生成 .docx。

用法:
    python generate_docx.py <report.json> [--template template.docx] [--output output.docx]
"""
import json
import os
import re
import sys

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_cell(cell, text, font_name="微软雅黑", font_size=10):
    """填入单元格文字，清除旧内容。"""
    # 清空旧 runs
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    # 清空多余段落
    for p2 in cell.paragraphs[1:]:
        for r in p2.runs:
            r.text = ""


def fill_table(table, rows_data, col_count):
    """用 rows_data 填充表格（第一行是表头，第二行是模板行）。"""
    template_row = table.rows[1]
    for i, row_data in enumerate(rows_data):
        row = template_row if i == 0 else table.add_row()
        for j in range(col_count):
            for p in row.cells[j].paragraphs:
                for r in p.runs:
                    r.text = ""
        for j, text in enumerate(row_data[:col_count]):
            set_cell(row.cells[j], str(text))


def clean_text(text):
    """去掉 @mention 和零宽空格。"""
    text = re.sub(r'@\S+\s*', '', text)
    text = re.sub(r'​', '', text)
    return text.strip()


def parse_args():
    args = {"report": None, "template": None, "output": None}
    argv = sys.argv[1:]
    i = 0
    while i < len(argv):
        if argv[i] == "--template" and i + 1 < len(argv):
            args["template"] = argv[i + 1]; i += 2
        elif argv[i] == "--output" and i + 1 < len(argv):
            args["output"] = argv[i + 1]; i += 2
        elif not argv[i].startswith("--") and args["report"] is None:
            args["report"] = argv[i]; i += 1
        else:
            i += 1
    return args


def generate(report_path, template_path=None, output_path=None):
    with open(report_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # 模板路径
    if not template_path:
        base = os.path.dirname(os.path.abspath(__file__))
        template_path = os.path.join(base, "..", "..", "..", "..", "XX项目周报-TB版.docx")
    template_path = os.path.abspath(template_path)

    # 输出路径
    if not output_path:
        out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "..", "output")
        os.makedirs(out_dir, exist_ok=True)
        output_path = os.path.join(out_dir, f"{data['project_name']}-项目周报-{data['week_start']}.docx")
    output_path = os.path.abspath(output_path)

    doc = Document(template_path)

    project_name = data["project_name"]
    week_num = data["week_num"].replace("2026年", "2026").replace("周", "周")

    # ── 替换占位符 ──
    for p in doc.paragraphs:
        full = "".join(r.text for r in p.runs if r.text)
        if not full.strip():
            continue
        new_text = full.replace("XX项目", project_name)
        new_text = new_text.replace("2026第XX周", week_num)
        new_text = new_text.replace("XX", project_name)
        new_text = new_text.replace(f"2026第{project_name}周", week_num)
        if p.runs:
            p.runs[0].text = new_text
            for r in p.runs[1:]:
                r.text = ""

    # ── 更新总数 ──
    section_order = ["本周已完成任务", "延期任务", "本周任务进展更新的详情", "截止时间为下周完成的任务"]
    counts = [len(data["completed"]), len(data["delayed"]), len(data["progress"]), len(data["next_week"])]
    idx = 0
    for p in doc.paragraphs:
        if p.text.strip() == "总数：" and idx < len(section_order):
            for r in p.runs:
                if "总数：" in r.text:
                    r.text = f"总数：{counts[idx]}项"
                    break
            idx += 1

    # ── 表格 0: 本周已完成 ──
    fill_table(doc.tables[0],
               [[t["title"], t["executor"], t["date"]] for t in data["completed"]], 3)

    # ── 表格 1: 延期任务 ──
    fill_table(doc.tables[1],
               [[t["title"], t["executor"], t["duedate"], f"{t['delay_days']}天"] for t in data["delayed"]], 4)

    # ── 表格 2: 进展更新 ──
    fill_table(doc.tables[2],
               [[t["title"], t["executor"], t["duedate"], clean_text(t["detail"]), ""] for t in data["progress"]], 5)

    # ── 表格 3: 下周截止 ──
    fill_table(doc.tables[3],
               [[t["title"], t["executor"], t["duedate"]] for t in data["next_week"]], 3)

    doc.save(output_path)
    print(f"Saved to {output_path}")
    return output_path


if __name__ == "__main__":
    args = parse_args()
    if not args["report"]:
        print("Usage: python generate_docx.py <report.json> [--template template.docx] [--output output.docx]")
        sys.exit(1)
    generate(args["report"], args.get("template"), args.get("output"))
